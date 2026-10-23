import copy
from pathlib import Path

import openpyxl

# -------------------------------------------------------------
# CONFIGURATION -- a dedicated, never-opened-in-Excel reference
# file that holds the one known-good pivot definition for
# 'Resumen - Contrato' / 'Resumen - Todo'. The live
# output_template.xlsx no longer carries baked-in pivots at all:
# every run injects a fresh clone from this reference instead,
# so a pivot degrading from Excel-side interaction (which has
# happened twice) can never propagate into a generated output.
# -------------------------------------------------------------
PIVOT_REFERENCE_PATH = Path("data/template/pivot_reference.xlsx")


# -------------------------------------------------------------
# BUILD: replace whatever pivot(s) exist (usually none) on each
# sheet of the output workbook with a fresh clone of the
# reference's pivot for that same sheet name, based on the
# output's own 'Data' sheet / Table1. Both 'Resumen - Contrato'
# and 'Resumen - Todo' clones share ONE cache object (mirroring
# how real Excel builds them) -- openpyxl loads each pivot's
# cache as a separate Python object even when they share one
# cacheId on disk, so cloning them independently would produce
# two caches both claiming the same cacheId, which Excel rejects.
# refreshOnLoad=True makes Excel compute the real pivot values
# against the new Data on open; includeNewItemsInFilter=True on
# page fields (e.g. Contract) ensures a different vendor's
# genuinely new values aren't excluded by a stale filter list.
# -------------------------------------------------------------
def rebuild_pivots_from_reference(workbook_path: Path, reference_path: Path = PIVOT_REFERENCE_PATH) -> None:
    if not reference_path.exists():
        raise FileNotFoundError(f"Pivot reference file not found: {reference_path}")

    ref_wb = openpyxl.load_workbook(reference_path)
    wb = openpyxl.load_workbook(workbook_path)

    ref_items = [
        (sheet_name, pivot)
        for sheet_name in ref_wb.sheetnames
        for pivot in getattr(ref_wb[sheet_name], "_pivots", [])
    ]
    if not ref_items:
        raise ValueError(f"No pivot tables found in reference file: {reference_path}")

    shared_cache = copy.deepcopy(ref_items[0][1].cache)
    shared_cache.refreshOnLoad = True

    for sheet_name, ref_pivot in ref_items:
        if sheet_name not in wb.sheetnames:
            continue
        target_ws = wb[sheet_name]
        target_ws._pivots = []

        cloned = copy.deepcopy(ref_pivot)
        cloned.cacheId = 1
        cloned.cache = shared_cache
        for field in cloned.pivotFields:
            if field.axis == "axisPage":
                field.includeNewItemsInFilter = True
                field.hideNewItems = False

        target_ws.add_pivot(cloned)

    wb.save(workbook_path)
