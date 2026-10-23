from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils import get_column_letter

AS_OF_DATE_FORMAT = "%d%b%Y"

# -------------------------------------------------------------
# CONFIGURATION
# -------------------------------------------------------------
DATA_SHEET_NAME = "Data"
PENDIENTES_SHEET_NAME = "Pendientes - Contrato"

DATE_NUMBER_FORMAT = "dd-mmm-yy"
DATE_COLUMNS = [
    "PO Document Date - Key",
    "Date : Calculated Delivery - Key",
    "Date: First GR Posting - Key",
    "Date: PO Statistical Delv - Key",
    "Date : PO Change - Key",
]

# Both sheets use shorter header names for these two source columns.
HEADER_NAME_OVERRIDES = {
    "Asset - Region (Text)": "Region",
    "Asset - Long Text": "Asset",
}

# Source columns the SEMANA/MES/DELIVERY/DEADLINE/PERFORMANCE/AÑO formulas read from.
COL_CALC_DELIVERY = "Date : Calculated Delivery - Key"
COL_GR_POSTING = "Date: First GR Posting - Key"
COL_DELIVERY_COMPLETION = "Indicator : Delivery Completion - Key"

# The six calculated columns appended after 'DIFOT%' in both sheets.
COL_SEMANA = "SEMANA"
COL_MES = "MES"
COL_DELIVERY = "DELIVERY"
COL_DEADLINE = "DEADLINE"
COL_PERFORMANCE = "PERFORMANCE"
COL_ANO = "AÑO"

WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
GRAY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


# -------------------------------------------------------------
# HELPER: look up a required header's column index, failing
# clearly if the sheet doesn't have it
# -------------------------------------------------------------
def _require_column(header_col_map: dict, sheet_name: str, header_name: str) -> int:
    col_idx = header_col_map.get(header_name)
    if col_idx is None:
        raise ValueError(f"'{sheet_name}' sheet is missing expected header: {header_name}")
    return col_idx


# -------------------------------------------------------------
# POPULATE: clear existing records (keep the header row), write
# the dataframe in their place, inject the calculated-column
# formulas, and apply zebra-stripe styling. Shared by every
# sheet we write to ('Data', 'Pendientes - Contrato', ...).
# -------------------------------------------------------------
def _populate_sheet(ws, sheet_name: str, df: pd.DataFrame, as_of_date: str) -> None:
    # 1. Delete existing data rows, keeping row 1 headers intact
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)

    # 2. Map each header name (row 1) to its column index
    header_col_map = {
        ws.cell(row=1, column=c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value is not None
    }

    # 3. Write the dataframe rows, aligning by header name
    for row_offset, row in enumerate(df.itertuples(index=False), start=2):
        for col_name, value in zip(df.columns, row):
            header_name = HEADER_NAME_OVERRIDES.get(col_name, col_name)
            col_idx = header_col_map.get(header_name)
            if col_idx is None:
                continue
            ws.cell(row=row_offset, column=col_idx, value=value)

    # 4. Resize the sheet's structured Table (if any) to match the new
    #    row count. Left stale, the Table keeps its old boundary and
    #    Excel auto-fills its calculated columns (SEMANA/MES/DELIVERY/...)
    #    down into the leftover blank rows, feeding bogus blank-date
    #    results into anything sourced from this table (e.g. pivot
    #    tables), which can blank out or break their grouping entirely.
    last_row = len(df) + 1
    last_col_letter = get_column_letter(ws.max_column)
    new_ref = f"A1:{last_col_letter}{last_row}"
    for table in ws.tables.values():
        table.ref = new_ref
        if table.autoFilter is not None:
            table.autoFilter.ref = new_ref

    # 5. Apply date formatting to the date columns, skipping "#" placeholders
    for col_name in DATE_COLUMNS:
        col_idx = header_col_map.get(col_name)
        if col_idx is None:
            continue
        for row_idx in range(2, last_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value != "#":
                cell.number_format = DATE_NUMBER_FORMAT

    # 6. Inject formulas into the six calculated columns appended after 'DIFOT%'
    calc_del_letter = get_column_letter(_require_column(header_col_map, sheet_name, COL_CALC_DELIVERY))
    gr_post_letter = get_column_letter(_require_column(header_col_map, sheet_name, COL_GR_POSTING))
    del_comp_letter = get_column_letter(_require_column(header_col_map, sheet_name, COL_DELIVERY_COMPLETION))

    semana_idx = _require_column(header_col_map, sheet_name, COL_SEMANA)
    mes_idx = _require_column(header_col_map, sheet_name, COL_MES)
    delivery_idx = _require_column(header_col_map, sheet_name, COL_DELIVERY)
    deadline_idx = _require_column(header_col_map, sheet_name, COL_DEADLINE)
    performance_idx = _require_column(header_col_map, sheet_name, COL_PERFORMANCE)
    ano_idx = _require_column(header_col_map, sheet_name, COL_ANO)

    delivery_letter = get_column_letter(delivery_idx)
    deadline_letter = get_column_letter(deadline_idx)

    # Frozen to the report's own date, not TODAY() -- otherwise DEADLINE/
    # PERFORMANCE (and everything derived from them: the pivot's OK/NOK
    # counts, the Resumen table, the chart) silently drift every time the
    # file is reopened later, instead of staying a snapshot of the report.
    as_of = datetime.strptime(as_of_date, AS_OF_DATE_FORMAT)
    as_of_literal = f"DATE({as_of.year},{as_of.month},{as_of.day})"

    for r in range(2, last_row + 1):
        ws.cell(row=r, column=semana_idx, value=f'="SEM " & WEEKNUM({calc_del_letter}{r})')
        ws.cell(row=r, column=mes_idx, value=f'=TEXT({calc_del_letter}{r}, "MMM-YYYY")')
        ws.cell(
            row=r,
            column=delivery_idx,
            value=(
                f'=IF(AND({gr_post_letter}{r}<>"#", {del_comp_letter}{r}="X"), "COMPLETA", '
                f'IF(AND({gr_post_letter}{r}<>"#", {del_comp_letter}{r}<>"X"), "PARCIAL", "PENDIENTE"))'
            ),
        )
        ws.cell(
            row=r,
            column=deadline_idx,
            value=(
                f'=IF(OR({delivery_letter}{r}="PENDIENTE", {delivery_letter}{r}="PARCIAL"), '
                f'{calc_del_letter}{r}-{as_of_literal}, {calc_del_letter}{r}-{gr_post_letter}{r})'
            ),
        )
        ws.cell(row=r, column=performance_idx, value=f'=IF({deadline_letter}{r}<0, "NOK", "OK")')
        ws.cell(row=r, column=ano_idx, value=f'=TEXT({calc_del_letter}{r}, "YYYY")')

    # 7. Zebra-stripe rows and add thin borders across every column
    for row_idx in range(2, last_row + 1):
        row_fill = GRAY_FILL if row_idx % 2 == 0 else WHITE_FILL
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.border = THIN_BORDER


# -------------------------------------------------------------
# WRITE: populate the 'Data' sheet with the vendor-filtered rows
# -------------------------------------------------------------
def write_data_sheet(workbook_path: Path, df: pd.DataFrame, as_of_date: str) -> None:
    wb = openpyxl.load_workbook(workbook_path)

    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{DATA_SHEET_NAME}' sheet not found in {workbook_path.name}")

    _populate_sheet(wb[DATA_SHEET_NAME], DATA_SHEET_NAME, df, as_of_date)
    wb.save(workbook_path)


# -------------------------------------------------------------
# WRITE: populate the 'Pendientes - Contrato' sheet with the
# pendientes-filtered rows (see transform.filter_pendientes)
# -------------------------------------------------------------
def write_pendientes_sheet(workbook_path: Path, df: pd.DataFrame, as_of_date: str) -> None:
    wb = openpyxl.load_workbook(workbook_path)

    if PENDIENTES_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{PENDIENTES_SHEET_NAME}' sheet not found in {workbook_path.name}")

    _populate_sheet(wb[PENDIENTES_SHEET_NAME], PENDIENTES_SHEET_NAME, df, as_of_date)
    wb.save(workbook_path)
