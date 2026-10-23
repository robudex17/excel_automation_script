"""
EXPERIMENT (isolated, disposable): can openpyxl build a working PivotTable
object from scratch -- not just edit one Excel already created -- well
enough that Excel accepts it and computes real numbers on open?

Does NOT touch src/excel_automation_script or data/template/output_template.xlsx.
Reads an already-generated output file (which has a populated 'Data' sheet /
Table1) purely as a data source, adds one new sheet holding the test pivot,
and saves as a separate file. Delete this whole experiments/ folder any time
to roll back with zero effect on the real pipeline.

Spec (from the user):
  Table: Table1 on 'Data' sheet
  Filter: Contract (Outline Agreement) No - Key
  Columns: Years(Date : Calculated Delivery - Key), Months(same field)
  Rows: PERFORMANCE, DELIVERY
  Values: Sum of No. of PO Schd Line
"""

from pathlib import Path
import shutil

import openpyxl
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, FieldGroup, RangePr, WorksheetSource
from openpyxl.pivot.table import DataField, Location, PageField, PivotField, RowColField, TableDefinition

SOURCE_FILE = Path("data/destination/CML COMERCIALIZADORA E IMPORTADORA - Semana 27 - 30Jun2026.xlsx")
OUTPUT_FILE = Path("experiments/pivot_test_output.xlsx")
PIVOT_SHEET_NAME = "Pivot Test"
DATA_TABLE_NAME = "Table1"
DATA_SHEET_NAME = "Data"

# 0-indexed positions of Table1's columns (see experiments/ investigation notes)
COL_NAMES = [
    "Region", "Asset", "Purch. Org. - Text", "Vendor(VN) - Key", "Vendor(VN) - Text",
    "IncoTerms - Key", "Inco term: Additional Information - Text", "Purchase Document Item No. - Key",
    "Purch. Group - Key", "Contract (Outline Agreement) No - Key", "Material - Key", "Material - Text",
    "Purchase Document Number - Key", "PO Item Quantity Order UOM", "Purch. Item Category - Text",
    "Indicator : Delivery Completion - Key", "PO Document Date - Key", "Date : Calculated Delivery - Key",
    "Date: First GR Posting - Key", "Date: PO Statistical Delv - Key", "Date : PO Change - Key",
    "No. of PO Schd Line", "Count of DIFOT", "DIFOT%", "SEMANA", "MES", "DELIVERY", "DEADLINE",
    "PERFORMANCE", "AÑO",
]
IDX_CONTRACT = COL_NAMES.index("Contract (Outline Agreement) No - Key")
IDX_CALC_DELIVERY = COL_NAMES.index("Date : Calculated Delivery - Key")
IDX_PO_SCHD_LINE = COL_NAMES.index("No. of PO Schd Line")
IDX_DELIVERY = COL_NAMES.index("DELIVERY")
IDX_PERFORMANCE = COL_NAMES.index("PERFORMANCE")

IDX_YEARS = len(COL_NAMES)      # 30, appended grouped field
IDX_MONTHS = len(COL_NAMES) + 1  # 31, appended grouped field


def build_cache_fields() -> list[CacheField]:
    fields = [CacheField(name=name) for name in COL_NAMES]

    years_group = FieldGroup(
        base=IDX_CALC_DELIVERY,
        rangePr=RangePr(autoStart=True, autoEnd=True, groupBy="years"),
    )
    months_group = FieldGroup(
        base=IDX_CALC_DELIVERY,
        rangePr=RangePr(autoStart=True, autoEnd=True, groupBy="months"),
    )
    fields.append(CacheField(name="Years", fieldGroup=years_group))
    fields.append(CacheField(name="Months", fieldGroup=months_group))
    return fields


def build_pivot_fields() -> list[PivotField]:
    total = len(COL_NAMES) + 2  # + Years, Months
    pivot_fields = []
    for i in range(total):
        if i == IDX_CONTRACT:
            pf = PivotField(axis="axisPage", showAll=False)
        elif i in (IDX_PERFORMANCE, IDX_DELIVERY):
            pf = PivotField(axis="axisRow", showAll=False)
        elif i in (IDX_YEARS, IDX_MONTHS):
            pf = PivotField(axis="axisCol", showAll=False)
        elif i == IDX_PO_SCHD_LINE:
            pf = PivotField(dataField=True, showAll=False)
        else:
            pf = PivotField(showAll=False)
        pivot_fields.append(pf)
    return pivot_fields


def main() -> None:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(f"Source file not found: {SOURCE_FILE}")

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_FILE, OUTPUT_FILE)

    wb = openpyxl.load_workbook(OUTPUT_FILE)

    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{DATA_SHEET_NAME}' sheet not found in {OUTPUT_FILE}")
    if PIVOT_SHEET_NAME in wb.sheetnames:
        del wb[PIVOT_SHEET_NAME]
    pivot_ws = wb.create_sheet(PIVOT_SHEET_NAME)

    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(name=DATA_TABLE_NAME)),
        cacheFields=build_cache_fields(),
        refreshOnLoad=True,
        recordCount=0,
    )

    pivot = TableDefinition(
        name="PivotFromScratch",
        cacheId=1,
        dataOnRows=False,
        dataCaption="Values",
        location=Location(ref="B3:C10", firstHeaderRow=1, firstDataRow=2, firstDataCol=1),
        pivotFields=build_pivot_fields(),
        rowFields=[RowColField(x=IDX_PERFORMANCE), RowColField(x=IDX_DELIVERY)],
        colFields=[RowColField(x=IDX_YEARS), RowColField(x=IDX_MONTHS)],
        pageFields=[PageField(fld=IDX_CONTRACT, item=0, hier=-1)],
        dataFields=[DataField(name="Sum of No. of PO Schd Line", fld=IDX_PO_SCHD_LINE, subtotal="sum", baseField=-1, baseItem=0)],
    )
    pivot.cache = cache

    pivot_ws.add_pivot(pivot)

    wb.save(OUTPUT_FILE)
    print(f"Saved: {OUTPUT_FILE}")

    # -----------------------------------------------------------------
    # Verify: reload and check the structure survives a round trip
    # -----------------------------------------------------------------
    wb2 = openpyxl.load_workbook(OUTPUT_FILE)
    reloaded = wb2[PIVOT_SHEET_NAME]._pivots
    print(f"Pivots on reload: {len(reloaded)}")
    if reloaded:
        p = reloaded[0]
        print(f"  name: {p.name}")
        print(f"  location.ref: {p.location.ref}")
        print(f"  rowFields: {[f.x for f in p.rowFields]}")
        print(f"  colFields: {[f.x for f in p.colFields]}")
        print(f"  pageFields: {[(f.fld, f.item) for f in p.pageFields]}")
        print(f"  dataFields: {[(f.name, f.fld, f.subtotal) for f in p.dataFields]}")
        print(f"  cache.refreshOnLoad: {p.cache.refreshOnLoad}")
        print(f"  cache.cacheFields count: {len(p.cache.cacheFields)}")


if __name__ == "__main__":
    main()
